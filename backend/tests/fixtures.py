from io import BytesIO

from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    PageBreak,
    Paragraph as RParagraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

_STYLES = getSampleStyleSheet()


def _cell(value):
    text = str(value)
    if "\n" in text:
        return RParagraph(text.replace("\n", "<br/>"), _STYLES["BodyText"])
    return text


def make_pdf(text: str) -> bytes:
    content = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode())
        out.write(obj)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return out.getvalue()


def build_pdf(
    blocks: list[tuple],
    *,
    header: str | None = None,
    footer: str | None = None,
    repeat_table_header: bool = True,
) -> bytes:
    """Build a PDF from a small story DSL.

    Blocks: ("title", text) | ("heading", text) | ("text", text)
            | ("table", header, rows, grid) | ("pagebreak",) | ("spacer", n)
    """
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=A4, title="test")

    def on_page(canvas, doc):
        canvas.saveState()
        if header:
            canvas.setFont("Helvetica", 9)
            canvas.drawString(40, A4[1] - 30, header)
        if footer:
            canvas.setFont("Helvetica", 9)
            canvas.drawString(40, 30, footer)
        canvas.restoreState()

    story = []
    for block in blocks:
        kind = block[0]
        if kind == "title":
            story.append(RParagraph(block[1], _STYLES["Title"]))
            story.append(Spacer(1, 10))
        elif kind == "heading":
            story.append(RParagraph(block[1], _STYLES["Heading2"]))
            story.append(Spacer(1, 6))
        elif kind == "text":
            story.append(RParagraph(block[1], _STYLES["BodyText"]))
            story.append(Spacer(1, 6))
        elif kind == "table":
            _, header_row, rows, grid = block
            data = [
                [_cell(value) for value in row]
                for row in ([header_row] if header_row else []) + rows
            ]
            table = Table(data, repeatRows=1 if repeat_table_header and header_row else 0)
            style = [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            if grid:
                style.append(("GRID", (0, 0), (-1, -1), 0.5, colors.black))
            table.setStyle(TableStyle(style))
            story.append(table)
            story.append(Spacer(1, 6))
        elif kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "spacer":
            story.append(Spacer(1, block[1]))

    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buffer.getvalue()


def make_docx(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_document(blocks: list[tuple]) -> bytes:
    document = DocxDocument()
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            document.add_heading(block[1], level=block[2] if len(block) > 2 else 1)
        elif kind == "paragraph":
            document.add_paragraph(block[1])
        elif kind == "list":
            for item in block[1]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            header, rows = block[1], block[2]
            table = document.add_table(rows=1 + len(rows), cols=len(header))
            for column, value in enumerate(header):
                table.rows[0].cells[column].text = value
            for row_index, row in enumerate(rows, start=1):
                for column, value in enumerate(row):
                    table.rows[row_index].cells[column].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_markdown(content: str) -> bytes:
    return content.encode("utf-8")
